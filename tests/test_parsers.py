"""Supported document parser contract tests."""

from __future__ import annotations

import hashlib
import io
from datetime import UTC, datetime

import pymupdf
import pytest
from docx import Document as DocxDocument

from shaper.domain import PermissionSnapshot, SourceDocument
from shaper.infrastructure.parsers import DocumentParseError, SupportedDocumentParser

ZERO_HASH = "0" * 64


def document_for(content: bytes, media_type: str) -> SourceDocument:
    """Create a source document matching content."""
    digest = hashlib.sha256(content).hexdigest()
    return SourceDocument(
        source_id="source-1",
        source_version=digest,
        content_hash=digest,
        tenant_id="tenant-1",
        collection_id="collection-1",
        title="Policy",
        media_type=media_type,
        observed_at=datetime(2026, 9, 9, tzinfo=UTC),
        permission=PermissionSnapshot(
            permission_hash=ZERO_HASH,
            captured_at=datetime(2026, 9, 9, tzinfo=UTC),
        ),
    )


def test_given_markdown_when_parsed_then_heading_path_and_order_are_preserved() -> None:
    # Arrange
    content = b"# Leave\n\nEmployees receive leave.\n\n## Exceptions\n\nContractors are excluded."

    # Act
    spans = SupportedDocumentParser().parse(document_for(content, "text/markdown"), content)

    # Assert
    assert [(span.heading_path, span.ordinal) for span in spans] == [
        (("Leave",), 0),
        (("Leave", "Exceptions"), 1),
    ]


def test_given_docx_when_parsed_then_paragraph_and_table_are_addressable() -> None:
    # Arrange
    document = DocxDocument()
    document.add_heading("Leave", level=1)
    document.add_paragraph("Employees receive leave.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Days"
    stream = io.BytesIO()
    document.save(stream)
    content = stream.getvalue()

    # Act
    spans = SupportedDocumentParser().parse(
        document_for(
            content,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
        content,
    )

    # Assert
    assert [span.text for span in spans] == ["Employees receive leave.", "Region", "Days"]


def test_given_docx_nested_content_when_parsed_then_no_policy_text_is_omitted() -> None:
    # Arrange
    document = DocxDocument()
    document.add_heading("Information security policy", level=1)
    outer = document.add_table(rows=1, cols=1)
    outer.cell(0, 0).paragraphs[0].text = "Policy owner: Security"
    nested = outer.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = (
        "Employees must report suspected incidents within 30 minutes "
        "and preserve all relevant evidence."
    )
    section = document.sections[0]
    section.header.paragraphs[0].text = "Controlled policy"
    stream = io.BytesIO()
    document.save(stream)
    content = stream.getvalue()

    # Act
    spans = SupportedDocumentParser().parse(
        document_for(
            content,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
        content,
    )

    # Assert
    extracted = "\n".join(span.text for span in spans)
    assert "Policy owner: Security" in extracted
    assert "Employees must report suspected incidents within 30 minutes" in extracted
    assert "preserve all relevant evidence" in extracted
    assert "Controlled policy" in extracted


def test_given_image_pdf_with_token_text_when_parsed_then_lossy_ingestion_is_rejected() -> None:
    # Arrange
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Policy")
    image = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 10, 10), 0)
    image.clear_with(255)
    page.insert_image(page.rect, stream=image.tobytes("png"))
    content = document.tobytes()
    document.close()

    # Act & Assert
    with pytest.raises(DocumentParseError, match="run OCR"):
        SupportedDocumentParser().parse(document_for(content, "application/pdf"), content)


def test_given_short_pdf_with_logo_when_parsed_then_readable_text_is_retained() -> None:
    # Arrange
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Approved policy statement.")
    image = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 10, 10), 0)
    image.clear_with(255)
    page.insert_image(pymupdf.Rect(72, 100, 120, 148), stream=image.tobytes("png"))
    content = document.tobytes()
    document.close()

    # Act
    spans = SupportedDocumentParser().parse(document_for(content, "application/pdf"), content)

    # Assert
    assert [span.text for span in spans] == ["Approved policy statement."]


def test_given_pdf_when_parsed_then_page_location_is_preserved() -> None:
    # Arrange
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Employees receive leave.")
    content = document.tobytes()
    document.close()

    # Act
    spans = SupportedDocumentParser().parse(document_for(content, "application/pdf"), content)

    # Assert
    assert spans[0].location["page"] == 1


def test_given_unsupported_media_type_when_parsed_then_diagnostic_is_explicit() -> None:
    # Arrange
    content = b"content"

    # Act & Assert
    with pytest.raises(DocumentParseError, match="unsupported media type"):
        SupportedDocumentParser().parse(document_for(content, "application/octet-stream"), content)

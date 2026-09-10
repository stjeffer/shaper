"""Safe parsers for digitally readable PDF, DOCX, Markdown, and text."""

from __future__ import annotations

import hashlib
import io
import unicodedata
import zipfile
from collections.abc import Iterable, Sequence

import pymupdf
from docx import Document as DocxDocument

from shaper.domain import SourceDocument, SourceSpan


class DocumentParseError(ValueError):
    """Raised when a supported document cannot produce safe readable spans."""


class SupportedDocumentParser:
    """Parse supported formats into ordered immutable spans."""

    def __init__(self, *, maximum_spans: int = 10_000) -> None:
        self._maximum_spans = maximum_spans

    def parse(self, document: SourceDocument, content: bytes) -> Sequence[SourceSpan]:
        """Parse content based on its validated media type."""
        try:
            if document.media_type == "application/pdf":
                blocks = self._pdf_blocks(content)
            elif (
                document.media_type
                == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ):
                blocks = self._docx_blocks(content)
            elif document.media_type in {"text/markdown", "text/plain"}:
                blocks = self._text_blocks(
                    content,
                    markdown=document.media_type == "text/markdown",
                )
            else:
                raise DocumentParseError(
                    f"Source {document.source_id!r} has unsupported media type "
                    f"{document.media_type!r}; use PDF, DOCX, Markdown, or plain text"
                )
            return self._spans(document, blocks)
        except DocumentParseError:
            raise
        except (OSError, RuntimeError, UnicodeError, ValueError, zipfile.BadZipFile) as error:
            raise DocumentParseError(
                f"Could not parse source {document.source_id!r} as {document.media_type!r}; "
                "verify that the file is digitally readable and not corrupt"
            ) from error

    def _spans(
        self,
        document: SourceDocument,
        blocks: Iterable[tuple[str, tuple[str, ...], dict[str, int | str]]],
    ) -> tuple[SourceSpan, ...]:
        spans: list[SourceSpan] = []
        for ordinal, (raw_text, headings, location) in enumerate(blocks):
            text = unicodedata.normalize("NFC", raw_text).strip()
            if not text:
                continue
            if len(spans) >= self._maximum_spans:
                raise DocumentParseError(
                    f"Source {document.source_id!r} exceeds the "
                    f"{self._maximum_spans}-span parser limit"
                )
            text_hash = hashlib.sha256(text.encode()).hexdigest()
            span_key = (
                f"{document.source_id}:{document.source_version}:{ordinal}:{text_hash}".encode()
            )
            spans.append(
                SourceSpan(
                    span_id=hashlib.sha256(span_key).hexdigest()[:32],
                    source_id=document.source_id,
                    source_version=document.source_version,
                    ordinal=ordinal,
                    text=text,
                    text_hash=text_hash,
                    heading_path=headings,
                    location=location,
                )
            )
        return tuple(spans)

    @staticmethod
    def _text_blocks(
        content: bytes,
        *,
        markdown: bool,
    ) -> Iterable[tuple[str, tuple[str, ...], dict[str, int | str]]]:
        text = content.decode("utf-8")
        headings: list[str] = []
        paragraph: list[str] = []
        block = 0
        for line in [*text.splitlines(), ""]:
            stripped = line.strip()
            if markdown and stripped.startswith("#"):
                if paragraph:
                    yield "\n".join(paragraph), tuple(headings), {"block": block}
                    paragraph.clear()
                    block += 1
                level = len(stripped) - len(stripped.lstrip("#"))
                heading = stripped[level:].strip()
                headings[level - 1 :] = [heading]
            elif stripped:
                paragraph.append(stripped)
            elif paragraph:
                yield "\n".join(paragraph), tuple(headings), {"block": block}
                paragraph.clear()
                block += 1

    @staticmethod
    def _docx_blocks(
        content: bytes,
    ) -> Iterable[tuple[str, tuple[str, ...], dict[str, int | str]]]:
        document = DocxDocument(io.BytesIO(content))
        headings: list[str] = []
        block = 0
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = paragraph.style.name if paragraph.style is not None else ""
            if style_name.startswith("Heading"):
                suffix = style_name.removeprefix("Heading").strip()
                level = int(suffix) if suffix.isdigit() else 1
                headings[level - 1 :] = [text]
                continue
            yield text, tuple(headings), {"block": block}
            block += 1
        for table_index, table in enumerate(document.tables):
            for row_index, row in enumerate(table.rows):
                text = " | ".join(cell.text.strip() for cell in row.cells)
                yield text, tuple(headings), {"table": table_index, "row": row_index}

    @staticmethod
    def _pdf_blocks(
        content: bytes,
    ) -> Iterable[tuple[str, tuple[str, ...], dict[str, int | str]]]:
        with pymupdf.open(stream=content, filetype="pdf") as document:
            for page_number, page in enumerate(document, start=1):
                for block_number, block in enumerate(page.get_text("blocks")):
                    text = str(block[4]).strip()
                    if text:
                        yield text, (), {"page": page_number, "block": block_number}
